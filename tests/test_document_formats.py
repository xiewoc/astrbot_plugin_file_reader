"""Real-document parser tests with locally generated fixtures.

Run after installing requirements. No real user document or account is required.
Fixtures and test-only dependencies are confined to the working directory.
"""

from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

WORKSPACE = Path(__file__).resolve().parents[1]
LOCAL_DEPENDENCIES = WORKSPACE / ".test-deps"
if LOCAL_DEPENDENCIES.is_dir():
    sys.path.insert(0, str(LOCAL_DEPENDENCIES))


def _create_pdf(path: Path, text: str | None) -> None:
    stream = (
        f"BT /F1 14 Tf 50 700 Td ({text}) Tj ET".encode("ascii") if text else b"q Q"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{i} 0 obj\n".encode())
        result.extend(obj)
        result.extend(b"\nendobj\n")
    startxref = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    result.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(b"trailer\n")
    result.extend(
        f"<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{startxref}\n%%EOF\n".encode()
    )
    path.write_bytes(result)


class RealDocuments(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Keep binary dependencies loaded outside the temporary AstrBot stubs.
        # Removing numpy from sys.modules while pandas retains its extension
        # objects causes a misleading 'cannot load module more than once' error.
        if importlib.util.find_spec("pandas") is not None:
            import pandas  # noqa: F401
        from test_plugin_hooks import _load_plugin

        cls.main = _load_plugin()

    def setUp(self):
        (WORKSPACE / ".local").mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=WORKSPACE / ".local")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.reader = self.main.FileReader(max_file_size=20 * 1024 * 1024)

    def _read(self, file: Path, marker: str):
        before = file.read_bytes()
        text = self.reader.read_file(file, max_length=40000)
        self.assertIn(marker, text)
        self.assertEqual(file.read_bytes(), before)
        return text

    def test_txt(self):
        f = self.root / "hello.txt"
        f.write_text("TXT_MARKER_971", encoding="utf-8")
        self._read(f, "TXT_MARKER_971")

    def test_markdown(self):
        f = self.root / "notes.md"
        f.write_text("## MD_MARKER_325", encoding="utf-8")
        self._read(f, "MD_MARKER_325")

    def test_csv(self):
        if not self.main.PANDAS_AVAILABLE:
            self.skipTest("pandas not installed")
        f = self.root / "report.csv"
        f.write_text("name,value\nCSV_MARKER_344,42\n", encoding="utf-8")
        self._read(f, "CSV_MARKER_344")

    def test_docx(self):
        if not self.main.DOCX2TXT_AVAILABLE or not self.main.PYTHON_DOCX_AVAILABLE:
            self.skipTest("Word parser dependencies not installed")
        from docx import Document

        f = self.root / "report.docx"
        doc = Document()
        doc.add_paragraph("DOCX_MARKER_729")
        doc.save(f)
        self._read(f, "DOCX_MARKER_729")

    def test_xlsx(self):
        if not self.main.PANDAS_AVAILABLE:
            self.skipTest("Excel parser dependencies not installed")
        from openpyxl import Workbook

        f = self.root / "report.xlsx"
        wb = Workbook()
        wb.active["A1"] = "XLSX_MARKER_281"
        wb.active["B1"] = 42
        wb.save(f)
        self._read(f, "XLSX_MARKER_281")

    def test_pptx(self):
        if not self.main.PYTHON_PPTX_AVAILABLE:
            self.skipTest("PowerPoint parser dependency not installed")
        from pptx import Presentation

        f = self.root / "slides.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(0, 0, 3000000, 500000)
        textbox.text = "PPTX_MARKER_818"
        prs.save(f)
        self._read(f, "PPTX_MARKER_818")

    def test_pdf_text(self):
        if not self.main.PDFMINER_AVAILABLE:
            self.skipTest("PDF parser not installed")
        f = self.root / "report.pdf"
        _create_pdf(f, "PDF_MARKER_681")
        self._read(f, "PDF_MARKER_681")

    def test_pdf_without_text(self):
        if not self.main.PDFMINER_AVAILABLE:
            self.skipTest("PDF parser not installed")
        f = self.root / "scanned.pdf"
        _create_pdf(f, None)
        with self.assertRaises(self.main.FileReaderError):
            self.reader.read_file(f, max_length=40000)

    def test_large_plain_text_truncated(self):
        f = self.root / "big.txt"
        f.write_text("A" * 60000 + "END_MARKER", encoding="utf-8")
        text = self.reader.read_file(f, max_length=40000)
        self.assertIn("内容截断", text)
        self.assertNotIn("END_MARKER", text)

    def test_oversized_office_expansion_rejected(self):
        f = self.root / "oversized.docx"
        with ZipFile(f, "w", compression=ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"A" * (2 * 1024 * 1024))
        self.assertLess(f.stat().st_size, 20 * 1024 * 1024)
        with self.assertRaisesRegex(self.main.FileReaderError, "压缩比例异常"):
            self.reader.read_file(f, max_length=40000)

    def test_corrupt_office_container_rejected(self):
        f = self.root / "corrupt.docx"
        f.write_bytes(b"PK\\x03\\x04invaliddata")
        with self.assertRaises(self.main.FileReaderError):
            self.reader.read_file(f, max_length=40000)

    def test_file_size_limit(self):
        f = self.root / "too_big.txt"
        f.write_text("X" * 1100, encoding="utf-8")
        small_reader = self.main.FileReader(max_file_size=1024)
        with self.assertRaises(self.main.FileTooLargeError):
            small_reader.read_file(f)

    def test_pdf_extraction_stops_at_character_budget(self):
        if not self.main.PDFMINER_AVAILABLE:
            self.skipTest("pdfminer not installed")
        f = self.root / "short.pdf"
        _create_pdf(f, "TINY")

        # Simulate a large single text element while preserving a tiny input.
        class FakeTextBox:
            def get_text(self):
                return "X" * 1000000

        with (
            patch.object(self.main, "LTTextContainer", FakeTextBox),
            patch.object(
                self.main, "pdf_extract_pages", return_value=[[FakeTextBox()]]
            ),
        ):
            text = self.reader.read_pdf_to_text(f, max_length=100)
        self.assertLessEqual(len(text), 101)

    def test_zip_guard_applies_to_renamed_office_archive(self):
        f = self.root / "renamed.dat"
        with ZipFile(f, "w", compression=ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"A" * (2 * 1024 * 1024))
        with (
            patch.object(self.reader, "detect_file_type", return_value="docx"),
            self.assertRaisesRegex(self.main.FileReaderError, "压缩比例异常"),
        ):
            self.reader.read_file(f, max_length=40000)

    def test_excel_all_sheets_failed_is_an_error(self):
        if not self.main.PANDAS_AVAILABLE:
            self.skipTest("pandas not installed")
        from openpyxl import Workbook

        f = self.root / "failure.xlsx"
        wb = Workbook()
        wb.active["A1"] = "data"
        wb.save(f)
        with (
            patch.object(
                self.main.pd, "read_excel", side_effect=ValueError("unreadable")
            ),
            self.assertRaises(self.main.FileReaderError),
        ):
            self.reader.read_file(f, max_length=40000)


if __name__ == "__main__":
    unittest.main()
